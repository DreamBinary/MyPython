from docx import Document
from openpyxl import Workbook
doc=Document(r"D:\Users\16477\PycharmProjects\pythonProject\word\猪猪.docx")
tab=doc.tables[0]
workbook=Workbook()
sheet=workbook.active
for i in range(len(tab.rows)):
    for j in range(len(tab.columns)):
        list1 = []
        list1.append(tab.cell(i,j).text)
        sheet.append(list1)
workbook.save(r"D:\Users\16477\PycharmProjects\pythonProject\word\w转e.xlsx")

























